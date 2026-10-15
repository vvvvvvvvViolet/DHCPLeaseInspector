"""Exports results to .xlsx: the status table and the software inventory."""

_STATUS_FILLS = {
    "Ready": "C8E6C9",          # green
    "Domain Issue": "FFF9C4",   # yellow
    "Patch Overdue": "FFF9C4",
    "Not in AD": "FFF9C4",
    "Offline": "FFCDD2",        # red
    "Error": "FFCDD2",
    "No DNS": "FFCDD2",
}

_FILL_GOOD = "C8E6C9"
_FILL_WARN = "FFF9C4"
_FILL_BAD = "FFCDD2"


def _append_text_row(sheet, values):
    """Append a row, keeping every value as literal text.

    openpyxl stores any string starting with "=" as a formula. Host names,
    user names and software titles all come from the scanned machines, so a
    crafted value like =HYPERLINK(...) would otherwise become a live formula
    in the auditor's workbook.
    """
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _fill_row(sheet, row_index, fill_hex):
    from openpyxl.styles import PatternFill

    fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    for cell in sheet[row_index]:
        cell.fill = fill


def _bold_header(sheet, headers):
    from openpyxl.styles import Font

    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"


def _autosize(sheet, cap=60):
    for column_cells in sheet.columns:
        lengths = [len(str(cell.value)) for cell in column_cells if cell.value is not None]
        width = max(lengths) if lengths else 0
        # Capped: one long error message shouldn't make a column page-wide.
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 10), cap)


def export_to_excel(path: str, rows: list[list[str]], headers: list[str]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DHCP Lease Inspector"
    _bold_header(sheet, headers)

    status_col = headers.index("Status") if "Status" in headers else None
    for row in rows:
        _append_text_row(sheet, row)
        if status_col is not None:
            fill_hex = _STATUS_FILLS.get(row[status_col])
            if fill_hex:
                _fill_row(sheet, sheet.max_row, fill_hex)

    _autosize(sheet)
    workbook.save(path)


def export_software(path: str, results: list[dict], policy=None) -> None:
    """Write a software inventory workbook.

    Sheets: "Software" (one row per computer x title, filterable), "Summary"
    (each title once, with how many computers have it and which versions),
    and "Scan Status" (one row per computer, so failures aren't silently
    missing from the inventory).

    With a policy that has rules, "Software" gains a Policy column flagging
    forbidden titles, "Scan Status" gains a verdict per computer, and a
    "Policy Violations" sheet lists every breach. Without rules the layout
    is unchanged.
    """
    from openpyxl import Workbook

    from . import policy as software_policy

    ordered = sorted(results, key=lambda r: r["computer"].casefold())
    use_policy = policy is not None and not policy.empty
    verdicts = {r["computer"]: software_policy.evaluate(r, policy) for r in ordered} if use_policy else {}
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Software"
    headers = ["Computer", "Software", "Version", "Publisher", "Install Date", "Scope"]
    _bold_header(sheet, headers + (["Policy"] if use_policy else []))
    for result in ordered:
        flagged = {}
        if use_policy:
            for rule, app in verdicts[result["computer"]]["forbidden"]:
                flagged.setdefault(id(app), []).append(rule)
        for app in result["apps"]:
            row = [result["computer"], app["name"], app["version"],
                   app["publisher"], app["install_date"], app["scope"]]
            if use_policy:
                rules = flagged.get(id(app))
                row.append("Forbidden: " + ", ".join(rules) if rules else "")
            _append_text_row(sheet, row)
            if use_policy and flagged.get(id(app)):
                _fill_row(sheet, sheet.max_row, _FILL_BAD)
    sheet.auto_filter.ref = sheet.dimensions
    _autosize(sheet)

    # One entry per title; versions and publishers vary between machines.
    titles: dict[str, dict] = {}
    for result in ordered:
        for app in result["apps"]:
            entry = titles.setdefault(app["name"].casefold(), {
                "name": app["name"], "publishers": set(), "versions": set(), "computers": set(),
            })
            if app["publisher"]:
                entry["publishers"].add(app["publisher"])
            if app["version"]:
                entry["versions"].add(app["version"])
            entry["computers"].add(result["computer"])

    summary = workbook.create_sheet("Summary")
    _bold_header(summary, ["Software", "Publisher", "Computers", "Versions"])
    for entry in sorted(titles.values(), key=lambda e: (-len(e["computers"]), e["name"].casefold())):
        _append_text_row(summary, [
            entry["name"],
            ", ".join(sorted(entry["publishers"])),
            len(entry["computers"]),
            ", ".join(sorted(entry["versions"])),
        ])
    summary.auto_filter.ref = summary.dimensions
    _autosize(summary)

    status = workbook.create_sheet("Scan Status")
    _bold_header(status, ["Computer", "Result", "Apps Found", "Error"]
                 + (["Policy"] if use_policy else []))
    for result in ordered:
        if not result["ok"]:
            verdict, fill_hex = "Failed", _FILL_BAD
        elif result.get("partial"):
            verdict, fill_hex = "Partial", _FILL_WARN
        else:
            verdict, fill_hex = "OK", _FILL_GOOD
        row = [result["computer"], verdict, len(result["apps"]), result.get("error") or ""]
        if use_policy:
            row.append(verdicts[result["computer"]]["status"])
        _append_text_row(status, row)
        _fill_row(status, status.max_row, fill_hex)
    status.auto_filter.ref = status.dimensions
    _autosize(status)

    if use_policy:
        breaches = workbook.create_sheet("Policy Violations")
        _bold_header(breaches, ["Computer", "Violation", "Rule", "Software Found", "Version"])
        for result in ordered:
            verdict = verdicts[result["computer"]]
            for rule, app in verdict["forbidden"]:
                _append_text_row(breaches, [result["computer"], "Forbidden software installed",
                                            rule, app["name"], app["version"]])
                _fill_row(breaches, breaches.max_row, _FILL_BAD)
            for rule in verdict["missing"]:
                _append_text_row(breaches, [result["computer"], "Required software missing",
                                            rule, "", ""])
                _fill_row(breaches, breaches.max_row, _FILL_WARN)
        breaches.auto_filter.ref = breaches.dimensions
        _autosize(breaches)

    workbook.save(path)
