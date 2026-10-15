"""End-to-end through the real window and the real worker threads (offscreen Qt)."""
import pytest
from conftest import days_ago, run_scan, wait_for
from PyQt5.QtCore import QItemSelectionModel
from PyQt5.QtWidgets import QTableWidgetItem

from dhcp_inspector import ad_utils, checks, config, dhcp_leases, software
from dhcp_inspector import main_window as mw
from dhcp_inspector import software_window as sww
from dhcp_inspector.main_window import (HEADERS, MainWindow, _CHANGE_COL, _DOMAIN_COL,
                                        _LAST_LOGON_COL, _MAC_COL, _OS_COL, _STATUS_COL, _USER_COL)


@pytest.fixture
def window(qapp):
    w = MainWindow()
    yield w
    w.close()


def ad_record(name, **kw):
    rec = {"name": name, "domain": "corp.local", "ad_os": "Windows 10 Pro", "enabled": True,
           "password_last_set": days_ago(3), "last_logon": days_ago(4)}
    return {**rec, **kw}


def load_ad(monkeypatch, window, *records):
    monkeypatch.setattr(ad_utils, "load_ad_computers", lambda *a, **k: list(records))
    window.on_load_ad()


def stub_scan(monkeypatch, probe=None, wmi=None):
    monkeypatch.setattr(checks, "resolve_and_ping", probe or (
        lambda t, timeout_ms=None: {"ip": "10.0.0.1", "dns_ok": True, "ping": True, "error": None}))
    monkeypatch.setattr(checks, "check_computer", wmi or (
        lambda n, timeout=None, credential=None: {"computer_name": n, "wmi_ok": True, "wmi_name": n,
                                                  "last_patch": days_ago(3), "error": None}))


def cell(window, row, col):
    return window.table.item(row, col).text()


def row_by(window, col, text):
    return next(r for r in range(window.table.rowCount()) if cell(window, r, col) == text)


def scan(qapp, window):
    run_scan(qapp, window, window.on_check_status)


# --- layout ---------------------------------------------------------------

def test_column_constants_match_the_headers_and_widths():
    assert len(HEADERS) == len(mw._COLUMN_WIDTHS)
    for const, header in [(_MAC_COL, "MAC"), (_USER_COL, "User Login"), (_DOMAIN_COL, "Domain"),
                          (_OS_COL, "OS Version"), (_LAST_LOGON_COL, "Last Logon"),
                          (_STATUS_COL, "Status"), (_CHANGE_COL, "Change")]:
        assert HEADERS[const] == header


def test_table_defaults_to_ascending_by_computer(window):
    # Regression: Qt's implicit default is column 0 descending.
    assert window.table.horizontalHeader().sortIndicatorOrder() == mw.Qt.AscendingOrder


# --- Load AD --------------------------------------------------------------

def test_load_ad_prefills_the_right_columns_and_leaves_no_holes(monkeypatch, window):
    load_ad(monkeypatch, window, ad_record("PC-1"))
    assert cell(window, 0, _DOMAIN_COL) == "corp.local" and cell(window, 0, _OS_COL) == "Windows 10 Pro"
    assert cell(window, 0, _LAST_LOGON_COL) != "-"
    assert all(window.table.item(0, c) is not None for c in range(len(HEADERS)))
    assert "Unscanned 1" in window.summary_label.text()


# --- scanning -------------------------------------------------------------

def test_scan_fills_status_and_summary(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("PC-1"), ad_record("PC-2"))
    stub_scan(monkeypatch)
    scan(qapp, window)
    assert [cell(window, r, _STATUS_COL) for r in range(2)] == ["Ready", "Ready"]
    assert "Ready 2" in window.summary_label.text() and window.check_status_btn.isEnabled()


def test_scan_runs_wmi_only_for_machines_that_answered_ping(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("UP"), ad_record("DOWN"))
    asked = []
    stub_scan(monkeypatch,
              probe=lambda t, timeout_ms=None: {"ip": "10.0.0.1", "dns_ok": True, "ping": t == "UP", "error": None},
              wmi=lambda n, timeout=None, credential=None: asked.append(n) or {
                  "computer_name": n, "wmi_ok": True, "last_patch": days_ago(3), "error": None})
    scan(qapp, window)
    assert asked == ["UP"]
    assert cell(window, row_by(window, 0, "DOWN"), _STATUS_COL) == "Offline"


def test_statuses_and_colours(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("OK"), ad_record("PATCH"),
            ad_record("STALE", password_last_set=days_ago(400)))
    stub_scan(monkeypatch, wmi=lambda n, timeout=None, credential=None: {
        "computer_name": n, "wmi_ok": True, "error": None,
        "last_patch": days_ago(300 if n == "PATCH" else 3)})
    scan(qapp, window)
    got = {n: cell(window, row_by(window, 0, n), _STATUS_COL) for n in ("OK", "PATCH", "STALE")}
    assert got == {"OK": "Ready", "PATCH": "Patch Overdue", "STALE": "Domain Issue"}
    bg = lambda n: window.table.item(row_by(window, 0, n), _STATUS_COL).background().color()
    assert bg("OK") == mw._COLOR_GOOD and bg("PATCH") == mw._COLOR_WARN and bg("STALE") == mw._COLOR_WARN


def test_mac_user_and_name_columns_from_wmi(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("PC-1"))
    stub_scan(monkeypatch,
              probe=lambda t, timeout_ms=None: {"ip": "10.0.0.1", "dns_ok": True, "ping": True,
                                                "arp_mac": "bb-bb-bb-bb-bb-bb", "error": None},
              wmi=lambda n, timeout=None, credential=None: {
                  "computer_name": n, "wmi_ok": True, "wmi_mac": "00:1a:2b:3c:4d:5e",
                  "user_login": "CORP\\somchai", "last_patch": days_ago(3), "error": None})
    scan(qapp, window)
    assert cell(window, 0, _MAC_COL) == "00-1A-2B-3C-4D-5E"       # WMI beats ARP
    assert cell(window, 0, _USER_COL) == "somchai"
    assert window.table.item(0, _USER_COL).toolTip() == "CORP\\somchai (domain CORP)"


def test_arp_mac_survives_a_refused_wmi_merge(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("PC-1"))
    stub_scan(monkeypatch,
              probe=lambda t, timeout_ms=None: {"ip": "10.0.0.1", "dns_ok": True, "ping": True,
                                                "arp_mac": "bb-bb-bb-bb-bb-bb", "error": None},
              wmi=lambda n, timeout=None, credential=None: {"computer_name": n, "wmi_ok": False,
                                                            "wmi_name": None, "error": "Access is denied."})
    scan(qapp, window)
    assert cell(window, 0, _MAC_COL) == "BB-BB-BB-BB-BB-BB" and "ARP" in window.table.item(0, _MAC_COL).toolTip()


def test_credentials_reach_the_wmi_call(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("PC-1"))
    seen = {}
    stub_scan(monkeypatch, wmi=lambda n, timeout=None, credential=None: seen.update(c=credential) or {
        "computer_name": n, "wmi_ok": True, "last_patch": days_ago(3), "error": None})
    window._credential = ("CORP\\svc", "pw")
    scan(qapp, window)
    assert seen["c"] == ("CORP\\svc", "pw")


# --- change tracking, re-check, filters ----------------------------------

def test_change_column_and_recheck_failed(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("PC-1"), ad_record("PC-2"))
    stub_scan(monkeypatch)
    scan(qapp, window)                                         # both Ready -> saved as the previous run
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: {
        "ip": "10.0.0.1", "dns_ok": True, "ping": t != "PC-2", "error": None})
    scan(qapp, window)
    assert cell(window, row_by(window, 0, "PC-2"), _CHANGE_COL) == "Ready → Offline"
    assert cell(window, row_by(window, 0, "PC-1"), _CHANGE_COL) == "-"

    probed = []
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: probed.append(t) or {
        "ip": "10.0.0.1", "dns_ok": True, "ping": True, "error": None})
    run_scan(qapp, window, window.on_recheck_failed)
    assert probed == ["PC-2"]
    assert cell(window, row_by(window, 0, "PC-2"), _STATUS_COL) == "Ready"
    assert cell(window, row_by(window, 0, "PC-2"), _CHANGE_COL) == "Offline → Ready"


def test_rows_land_on_the_right_machine_after_the_user_resorts(monkeypatch, qapp, window):
    # Row indices are positional; sorting mid-flight must not scramble results.
    load_ad(monkeypatch, window, ad_record("A"), ad_record("B"), ad_record("C"))
    stub_scan(monkeypatch)
    window.table.sortItems(0, mw.Qt.DescendingOrder)
    scan(qapp, window)
    for r in range(3):
        assert cell(window, r, _STATUS_COL) == "Ready"
        assert cell(window, r, 0) in "ABC"


def test_status_filter_and_live_only_compose(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("UP"), ad_record("DOWN"))
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: {
        "ip": "10.0.0.1", "dns_ok": True, "ping": t == "UP", "error": None})
    scan(qapp, window)
    up, down = row_by(window, 0, "UP"), row_by(window, 0, "DOWN")
    window.live_only.setChecked(True)
    assert not window.table.isRowHidden(up) and window.table.isRowHidden(down)
    window.live_only.setChecked(False)
    window.status_filter.setCurrentText("Offline")
    assert window.table.isRowHidden(up) and not window.table.isRowHidden(down)


# --- subnet ---------------------------------------------------------------

def test_load_subnet_rejects_an_oversized_range_without_freezing(monkeypatch, window):
    warned = []
    monkeypatch.setattr(mw.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    window.subnet_edit.setText("10.0.0.0/8")
    window.on_load_subnet()
    assert window.table.rowCount() == 0 and "host limit" in warned[0]


def test_subnet_scan_names_hosts_and_flags_ones_missing_from_ad(monkeypatch, qapp, window):
    load_ad(monkeypatch, window, ad_record("A-C14-PROV"))
    window.subnet_edit.setText("10.20.30.5-7")
    window.on_load_subnet()
    leases = {"10.20.30.6": {"name": "PRINTER-2", "mac": "00-11-22-33-44-55"}}
    config.save_settings(config.Settings(dhcp_server="dhcp01"))
    monkeypatch.setattr(dhcp_leases, "load_leases", lambda server: leases)
    stub_scan(monkeypatch,
              probe=lambda t, timeout_ms=None: {"ip": t, "dns_ok": True, "ping": True, "arp_mac": None,
                                                "resolved_name": None, "error": None},
              wmi=lambda n, timeout=None, credential=None: (
                  {"computer_name": n, "wmi_ok": True, "wmi_name": "A-C14-PROV", "last_patch": days_ago(3), "error": None}
                  if n == "10.20.30.5" else
                  {"computer_name": n, "wmi_ok": False, "wmi_name": None, "error": "Access is denied."}))
    scan(qapp, window)
    by_ip = {cell(window, r, 1): r for r in range(3)}
    assert cell(window, by_ip["10.20.30.5"], 0) == "A-C14-PROV"                  # named by WMI ...
    assert cell(window, by_ip["10.20.30.5"], _STATUS_COL) == "Ready"             # ... and matched to AD
    assert cell(window, by_ip["10.20.30.6"], 0) == "PRINTER-2"                    # named by DHCP lease
    assert cell(window, by_ip["10.20.30.6"], _MAC_COL) == "00-11-22-33-44-55"
    assert cell(window, by_ip["10.20.30.6"], _STATUS_COL) == "Not in AD"
    assert cell(window, by_ip["10.20.30.7"], 0) == "10.20.30.7"                  # nothing could name it


def test_subnet_without_ad_never_claims_not_in_ad(monkeypatch, qapp, window):
    window.subnet_edit.setText("10.20.30.1")
    window.on_load_subnet()
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: {
        "ip": t, "dns_ok": True, "ping": True, "arp_mac": None, "resolved_name": None, "error": None})
    scan(qapp, window)
    assert cell(window, 0, _STATUS_COL) == "Ready"


def test_a_failing_dhcp_lookup_degrades_to_a_warning(monkeypatch, qapp, window):
    window.subnet_edit.setText("10.20.30.1")
    window.on_load_subnet()
    config.save_settings(config.Settings(dhcp_server="dhcp01"))
    monkeypatch.setattr(dhcp_leases, "load_leases", lambda s: (_ for _ in ()).throw(RuntimeError("RSAT missing")))
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: {
        "ip": t, "dns_ok": True, "ping": True, "arp_mac": None, "resolved_name": None, "error": None})
    scan(qapp, window)
    assert "RSAT missing" in window.summary_label.text() and window.check_status_btn.isEnabled()


# --- cancel ---------------------------------------------------------------

def test_cancel_stops_the_scan_and_re_enables_the_buttons(monkeypatch, qapp, window):
    import time
    load_ad(monkeypatch, window, *[ad_record(f"PC-{i}") for i in range(40)])
    config.save_settings(config.Settings(max_parallel_ping=2))
    stub_scan(monkeypatch, probe=lambda t, timeout_ms=None: (time.sleep(0.05), {
        "ip": "10.0.0.1", "dns_ok": True, "ping": False, "error": None})[1])
    window.on_check_status()
    wait_for(qapp, lambda: window.cancel_btn.isEnabled())
    window.on_cancel()
    wait_for(qapp, lambda: not window._worker.isRunning())
    assert window.check_status_btn.isEnabled() and window.recheck_btn.isEnabled()


# --- settings dialog ------------------------------------------------------

def test_settings_dialog_round_trips_every_field(qapp, window):
    d = mw.SettingsDialog(window)
    d.stale_days.setValue(90); d.stale_patch_days.setValue(30); d.dhcp_server.setText("dhcp01")
    d.max_subnet_hosts.setValue(512); d.software_timeout.setValue(600)
    d.save()
    s = config.load_settings()
    assert (s.stale_password_days, s.stale_patch_days, s.dhcp_server, s.max_subnet_hosts, s.software_timeout_s) \
        == (90, 30, "dhcp01", 512, 600)


# --- software inventory ---------------------------------------------------

def mark_wmi(window, names_ok):
    window.table.setSortingEnabled(False)
    for name, ok in names_ok.items():
        r = window.table.rowCount(); window.table.insertRow(r)
        item = QTableWidgetItem(name); item.setData(mw._WMI_OK_ROLE, ok)
        window.table.setItem(r, 0, item)


def test_software_targets_default_to_rows_that_reached_wmi(window):
    mark_wmi(window, {"UP-1": True, "DOWN": False, "UP-2": True})
    assert window._software_targets() == (["UP-1", "UP-2"], "reachable over WMI")


def test_software_targets_prefer_visible_selected_rows(window):
    mark_wmi(window, {"PC-0": True, "PC-1": True, "PC-2": False})
    sm = window.table.selectionModel()
    for r in (0, 1, 2):
        sm.select(window.table.model().index(r, 0), QItemSelectionModel.Select | QItemSelectionModel.Rows)
    window.table.setRowHidden(0, True)                         # filtered out: must not be sent
    assert window._software_targets() == (["PC-1", "PC-2"], "selected")


def test_software_with_nothing_to_inventory_explains_itself(monkeypatch, window):
    shown = []
    monkeypatch.setattr(mw.QMessageBox, "information", lambda *a: shown.append(a[2]))
    window.on_software()
    assert shown and window._software_window is None


def test_software_asks_before_a_big_fleet_and_honours_no(monkeypatch, window):
    mark_wmi(window, {f"PC-{i}": True for i in range(mw._CONFIRM_SOFTWARE_OVER + 1)})
    monkeypatch.setattr(mw.QMessageBox, "question", lambda *a: mw.QMessageBox.No)
    window.on_software()
    assert window._software_window is None


def test_software_window_inventories_and_exports(monkeypatch, qapp, window, tmp_path):
    mark_wmi(window, {"A": True, "B": True})
    ok = {"ok": True, "partial": False, "error": None,
          "apps": [{"name": "Google Chrome", "version": "129", "publisher": "Google",
                    "install_date": "2026-09-01", "scope": "Machine"}]}
    monkeypatch.setattr(software, "list_software", lambda c, credential=None, timeout=None: (
        {**ok, "computer": c} if c == "A" else
        {"computer": c, "ok": False, "partial": False, "apps": [], "error": "Access is denied."}))
    window.on_software()
    sw_win = window._software_window
    wait_for(qapp, lambda: not sw_win.is_running())
    assert "OK 1" in sw_win.summary_label.text() and "Failed 1" in sw_win.summary_label.text()
    texts = [sw_win.table.item(r, 1).text() for r in range(sw_win.table.rowCount())]
    assert texts == ["Google Chrome", "(scan failed)"]            # the failure stays visible
    sw_win.filter_edit.setText("chrome")
    assert [sw_win.table.isRowHidden(r) for r in range(2)] == [False, True]
    out = tmp_path / "inv.xlsx"
    sw_win.export_to(str(out))
    from openpyxl import load_workbook
    assert load_workbook(out).sheetnames == ["Software", "Summary", "Scan Status"]
    sw_win.close()


def test_software_table_defaults_to_ascending(qapp):
    win = sww.SoftwareWindow(["A"], None)
    assert win.table.horizontalHeader().sortIndicatorOrder() == sww.Qt.AscendingOrder
    win.close()
