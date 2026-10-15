import zipfile

import pytest
from openpyxl import load_workbook

from dhcp_inspector import export_excel
from dhcp_inspector import policy as sp
from dhcp_inspector import software_window as sww
from dhcp_inspector.policy import Policy


def app(name, version="1.0"):
    return {"name": name, "version": version, "publisher": "Pub", "install_date": "2026-09-01", "scope": "Machine"}


def result(computer, apps, ok=True, partial=False, error=None):
    return {"computer": computer, "ok": ok, "partial": partial, "error": error, "apps": apps}


# --- rule parsing and matching --------------------------------------------

def test_parse_rules_drops_blanks_comments_and_duplicates_keeping_order():
    text = "  AnyDesk \n\n# remote tools\nTeamViewer\nanydesk\n   \nuTorrent"
    assert sp.parse_rules(text) == ["AnyDesk", "TeamViewer", "uTorrent"]
    assert sp.parse_rules("") == [] and sp.parse_rules(None) == []


@pytest.mark.parametrize("rule, name, expected", [
    ("teamviewer", "TeamViewer 15 Host", True),            # substring, any case
    ("TEAMVIEWER", "teamviewer", True),
    ("anydesk*", "AnyDesk", True),
    ("microsoft * 2019", "Microsoft SQL Server 2019", True),
    ("microsoft * 2019", "Microsoft SQL Server 2017", False),
    ("sophos", "Microsoft Defender", False),
    ("Visual C++ 2015 (x64)", "Microsoft Visual C++ 2015 (x64) Redistributable", True),  # literal ()+
    ("c++", "C Sharp", False),
    ("[a-z]", "abc", False),                                # brackets are not a character class
    ("a.c", "abc", False),                                  # "." is not a wildcard
    ("", "anything", True),                                 # blank rules never reach here (parse drops them)
])
def test_rule_matching(rule, name, expected):
    assert sp.matches(rule, name) is expected


# --- evaluation -----------------------------------------------------------

POLICY = Policy(forbidden=["AnyDesk", "uTorrent"], required=["Sophos"])


def test_compliant_machine():
    v = sp.evaluate(result("A", [app("Sophos Endpoint"), app("Google Chrome")]), POLICY)
    assert v == {"status": "Compliant", "forbidden": [], "missing": []}


def test_forbidden_software_is_a_violation_and_names_the_app():
    r = result("A", [app("Sophos Endpoint"), app("AnyDesk 8", "8.0.1")])
    v = sp.evaluate(r, POLICY)
    assert v["status"] == "Violation" and v["missing"] == []
    assert [(rule, a["name"]) for rule, a in v["forbidden"]] == [("AnyDesk", "AnyDesk 8")]


def test_missing_required_software_is_a_violation():
    v = sp.evaluate(result("A", [app("Google Chrome")]), POLICY)
    assert v["status"] == "Violation" and v["missing"] == ["Sophos"]


def test_failed_scan_is_unknown_not_missing():
    # We don't know what is installed, so we must not claim Sophos is absent.
    v = sp.evaluate(result("A", [], ok=False, error="Access is denied."), POLICY)
    assert v == {"status": "Unknown", "forbidden": [], "missing": []}


def test_partial_list_never_claims_missing_but_still_reports_forbidden_it_saw():
    partial = result("A", [app("Google Chrome")], partial=True)
    assert sp.evaluate(partial, POLICY) == {"status": "Unknown", "forbidden": [], "missing": []}
    seen = result("A", [app("uTorrent")], partial=True)
    v = sp.evaluate(seen, POLICY)
    assert v["status"] == "Violation" and v["missing"] == [] and len(v["forbidden"]) == 1


def test_forbidden_only_policy_on_a_complete_clean_list_is_compliant():
    assert sp.evaluate(result("A", [app("Chrome")]), Policy(forbidden=["AnyDesk"]))["status"] == "Compliant"


def test_one_app_matching_two_rules_is_reported_for_each():
    v = sp.evaluate(result("A", [app("AnyDesk Tool")]), Policy(forbidden=["anydesk", "desk"]))
    assert sorted(rule for rule, _ in v["forbidden"]) == ["anydesk", "desk"]


# --- storage --------------------------------------------------------------

def test_policy_round_trips_and_keeps_non_ascii():
    sp.save_policy(Policy(forbidden=["AnyDesk", "โปรแกรมต้องห้าม"], required=["Sophos"]))
    loaded = sp.load_policy()
    assert loaded.forbidden == ["AnyDesk", "โปรแกรมต้องห้าม"] and loaded.required == ["Sophos"]


def test_missing_or_corrupt_policy_file_means_no_rules():
    assert sp.load_policy().empty
    sp._policy_path().write_text("{nope", encoding="utf-8")
    assert sp.load_policy().empty
    sp._policy_path().write_text("[1, 2]", encoding="utf-8")     # valid JSON, wrong shape
    assert sp.load_policy().empty


def test_loading_normalises_a_hand_edited_file():
    sp._policy_path().write_text('{"forbidden": ["  AnyDesk ", "anydesk", ""], "required": []}', encoding="utf-8")
    assert sp.load_policy().forbidden == ["AnyDesk"]


# --- Excel ----------------------------------------------------------------

RESULTS = [
    result("PC-A", [app("Sophos Endpoint"), app("Google Chrome")]),
    result("PC-B", [app("Google Chrome"), app("AnyDesk 8", "8.0.1")]),
    result("PC-C", [app("Sophos Endpoint"), app("uTorrent", "3.5")]),
    result("PC-D", [], ok=False, error="Access is denied."),
]


def export(tmp_path, results=RESULTS, policy=POLICY):
    path = tmp_path / "sw.xlsx"
    export_excel.export_software(str(path), results, policy)
    return path, load_workbook(path)


def test_export_without_rules_keeps_the_original_layout(tmp_path):
    for policy in (None, Policy()):
        _, wb = export(tmp_path, policy=policy)
        assert wb.sheetnames == ["Software", "Summary", "Scan Status"]
        assert [c.value for c in wb["Software"][1]] == ["Computer", "Software", "Version", "Publisher", "Install Date", "Scope"]
        assert [c.value for c in wb["Scan Status"][1]] == ["Computer", "Result", "Apps Found", "Error"]


def test_export_with_rules_adds_policy_column_verdicts_and_a_violations_sheet(tmp_path):
    _, wb = export(tmp_path)
    assert wb.sheetnames == ["Software", "Summary", "Scan Status", "Policy Violations"]

    sw = wb["Software"]
    assert sw["G1"].value == "Policy"
    flagged = {(r[0].value, r[1].value): r[6].value for r in sw.iter_rows(min_row=2)}
    assert flagged[("PC-B", "AnyDesk 8")] == "Forbidden: AnyDesk"
    assert flagged[("PC-A", "Google Chrome")] in ("", None)
    row = next(r for r in sw.iter_rows(min_row=2) if r[1].value == "AnyDesk 8")
    assert row[1].fill.start_color.rgb[-6:] == "FFCDD2"

    verdicts = {r[0].value: r[4].value for r in wb["Scan Status"].iter_rows(min_row=2)}
    assert verdicts == {"PC-A": "Compliant", "PC-B": "Violation", "PC-C": "Violation", "PC-D": "Unknown"}


def test_violations_sheet_lists_each_breach(tmp_path):
    _, wb = export(tmp_path)
    rows = [tuple(c.value for c in r) for r in wb["Policy Violations"].iter_rows(min_row=2)]
    assert rows == [
        ("PC-B", "Forbidden software installed", "AnyDesk", "AnyDesk 8", "8.0.1"),
        ("PC-B", "Required software missing", "Sophos", None, None),
        ("PC-C", "Forbidden software installed", "uTorrent", "uTorrent", "3.5"),
    ]


def test_a_failed_machine_is_never_listed_as_missing_software(tmp_path):
    _, wb = export(tmp_path)
    assert "PC-D" not in {r[0].value for r in wb["Policy Violations"].iter_rows(min_row=2)}


def test_policy_text_cannot_inject_formulas_into_the_workbook(tmp_path):
    evil = '=HYPERLINK("http://evil.example","x")'
    path, wb = export(tmp_path, results=[result("PC-A", [app("Chrome")])], policy=Policy(required=[evil]))
    zf = zipfile.ZipFile(path)
    for member in zf.namelist():
        if member.startswith("xl/worksheets/sheet"):
            assert "<f>" not in zf.read(member).decode(), member
    assert wb["Policy Violations"]["C2"].value == evil


# --- GUI ------------------------------------------------------------------

def test_policy_dialog_round_trips_text(qapp):
    dialog = sww.PolicyDialog(Policy(forbidden=["AnyDesk"], required=["Sophos"]))
    assert dialog.forbidden_edit.toPlainText() == "AnyDesk"
    dialog.forbidden_edit.setPlainText("AnyDesk\n# note\n\nuTorrent\nanydesk")
    dialog.required_edit.setPlainText("Sophos")
    assert dialog.policy() == Policy(forbidden=["AnyDesk", "uTorrent"], required=["Sophos"])


def test_edit_policy_saves_on_ok_and_leaves_it_alone_on_cancel(monkeypatch, qapp):
    monkeypatch.setattr(sww.PolicyDialog, "exec_", lambda self: sww.QDialog.Accepted)
    sp.save_policy(Policy(forbidden=["Old"]))
    monkeypatch.setattr(sww.PolicyDialog, "policy", lambda self: Policy(forbidden=["New"], required=["Sophos"]))
    assert sww.edit_policy() == Policy(forbidden=["New"], required=["Sophos"])
    assert sp.load_policy().forbidden == ["New"]

    monkeypatch.setattr(sww.PolicyDialog, "exec_", lambda self: sww.QDialog.Rejected)
    monkeypatch.setattr(sww.PolicyDialog, "policy", lambda self: Policy(forbidden=["Ignored"]))
    assert sww.edit_policy() is None
    assert sp.load_policy().forbidden == ["New"]


def make_window(qapp):
    win = sww.SoftwareWindow(["PC-A", "PC-B", "PC-D"], None)
    for r in (RESULTS[0], RESULTS[1], RESULTS[3]):
        win.on_result(r)
    return win


def cells(win, col):
    return [win.table.item(r, col).text() for r in range(win.table.rowCount())]


def test_window_without_rules_shows_no_policy_noise(qapp):
    win = make_window(qapp)
    assert set(cells(win, sww._POLICY_COL)) == {""}
    assert "Violations" not in win.summary_label.text()
    win.close()


def test_window_flags_forbidden_and_missing_and_summarises(qapp):
    sp.save_policy(POLICY)
    win = make_window(qapp)
    by_software = dict(zip(cells(win, 1), cells(win, sww._POLICY_COL)))
    assert by_software["AnyDesk 8"] == "Forbidden: AnyDesk"
    assert by_software["(missing required: Sophos)"] == "Missing required"
    assert by_software["(scan failed)"] == "Unknown"
    text = win.summary_label.text()
    assert "Violations 1" in text and "Compliant 1" in text and "Unknown 1" in text
    colours = {win.table.item(r, 1).text(): win.table.item(r, 1).background().color() for r in range(win.table.rowCount())}
    assert colours["AnyDesk 8"] == sww._COLOR_FORBIDDEN
    assert colours["(missing required: Sophos)"] == sww._COLOR_MISSING
    win.close()


def test_editing_the_policy_redraws_existing_results_and_export_follows(monkeypatch, qapp, tmp_path):
    win = make_window(qapp)
    before = win.table.rowCount()
    monkeypatch.setattr(sww, "edit_policy", lambda parent=None: Policy(forbidden=["Chrome"]))
    win.on_edit_policy()
    assert "Forbidden: Chrome" in cells(win, sww._POLICY_COL)
    assert win.table.rowCount() == before                       # forbidden adds a flag, not rows
    assert "Violations 2" in win.summary_label.text()           # PC-A and PC-B both have Chrome

    out = tmp_path / "x.xlsx"
    win.export_to(str(out))
    assert "Policy Violations" in load_workbook(out).sheetnames

    monkeypatch.setattr(sww, "edit_policy", lambda parent=None: None)   # cancelled: nothing changes
    win.on_edit_policy()
    assert "Violations 2" in win.summary_label.text()
    win.close()


def test_filter_can_find_violations(qapp):
    sp.save_policy(POLICY)
    win = make_window(qapp)
    win.filter_edit.setText("forbidden")
    shown = [win.table.item(r, 1).text() for r in range(win.table.rowCount()) if not win.table.isRowHidden(r)]
    assert shown == ["AnyDesk 8"]
    win.close()
