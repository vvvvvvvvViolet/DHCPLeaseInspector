import zipfile

import pytest
from openpyxl import load_workbook

from dhcp_inspector import config, export_excel, history


# --- settings -------------------------------------------------------------

def test_settings_defaults_and_round_trip():
    s = config.load_settings()
    assert s.stale_password_days == 45 and s.wmi_only_ping_ok is True and s.dhcp_server == ""
    s.stale_password_days, s.dhcp_server = 60, "dhcp01"
    config.save_settings(s)
    config._settings = None                      # force a re-read from disk
    again = config.get_settings()
    assert again.stale_password_days == 60 and again.dhcp_server == "dhcp01"


def test_settings_survive_corrupt_or_old_files():
    config._settings_path().write_text("{not json", encoding="utf-8")
    assert config.load_settings() == config.Settings()
    # a file written by an older version, missing newer keys
    config._settings_path().write_text('{"stale_password_days": 10}', encoding="utf-8")
    s = config.load_settings()
    assert s.stale_password_days == 10 and s.max_subnet_hosts == 4096


# --- history --------------------------------------------------------------

def test_history_records_and_returns_the_last_run():
    assert history.last_run() is None
    history.save_run({"PC-1": "Ready"})
    history.save_run({"PC-1": "Offline"})
    last = history.last_run()
    assert last["statuses"] == {"PC-1": "Offline"} and "timestamp" in last


def test_history_keeps_only_the_most_recent_runs():
    for _ in range(40):
        history.save_run({"PC-1": "Ready"})
    assert len(history.load_runs()) == 30


def test_history_ignores_a_corrupt_file():
    history._runs_path().write_text("oops", encoding="utf-8")
    assert history.load_runs() == [] and history.last_run() is None


# --- status export --------------------------------------------------------

def test_status_export_colours_by_status(tmp_path):
    path = tmp_path / "out.xlsx"
    headers = ["Computer", "Status"]
    export_excel.export_to_excel(str(path), [["A", "Ready"], ["B", "Patch Overdue"], ["C", "Offline"]], headers)
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == headers and ws.freeze_panes == "A2"
    fills = {ws.cell(r, 1).value: ws.cell(r, 2).fill.start_color.rgb[-6:] for r in (2, 3, 4)}
    assert fills == {"A": "C8E6C9", "B": "FFF9C4", "C": "FFCDD2"}


@pytest.mark.parametrize("payload", ['=HYPERLINK("http://evil.example","x")', '=cmd|"/c calc"!A1'])
def test_status_export_stores_formula_lookalikes_as_text(tmp_path, payload):
    path = tmp_path / "out.xlsx"
    export_excel.export_to_excel(str(path), [[payload, "Ready"]], ["Computer", "Status"])
    assert "<f>" not in zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    assert load_workbook(path).active["A2"].value == payload


# --- software export ------------------------------------------------------

def sw(name, version, publisher="Pub", scope="Machine"):
    return {"name": name, "version": version, "publisher": publisher, "install_date": "2026-09-01", "scope": scope}


RESULTS = [
    {"computer": "PC-B", "ok": True, "partial": False, "error": None,
     "apps": [sw("Google Chrome", "129.0"), sw("7-Zip", "24.08", "Igor Pavlov"), sw('=HYPERLINK("http://evil.example","x")', "1")]},
    {"computer": "PC-A", "ok": True, "partial": False, "error": None,
     "apps": [sw("Google Chrome", "128.0"), sw("Zoom", "6.2", "Zoom", "Per-user")]},
    {"computer": "PC-C", "ok": True, "partial": True, "error": "RPC unavailable", "apps": [sw("Google Chrome", "129.0")]},
    {"computer": "PRINTER-2", "ok": False, "partial": False, "error": "Access is denied.", "apps": []},
]


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "sw.xlsx"
    export_excel.export_software(str(path), RESULTS)
    return path, load_workbook(path)


def test_software_workbook_has_three_sheets(workbook):
    _, wb = workbook
    assert wb.sheetnames == ["Software", "Summary", "Scan Status"]


def test_software_sheet_is_sorted_filtered_and_frozen(workbook):
    _, wb = workbook
    ws = wb["Software"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Computer", "Software", "Version", "Publisher", "Install Date", "Scope")
    assert [r[0] for r in rows[1:]] == ["PC-A", "PC-A", "PC-B", "PC-B", "PC-B", "PC-C"]
    assert ws.auto_filter.ref and ws.freeze_panes == "A2"


def test_summary_counts_computers_and_versions_most_widespread_first(workbook):
    _, wb = workbook
    rows = list(wb["Summary"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "Google Chrome" and rows[0][2] == 3 and rows[0][3] == "128.0, 129.0"


def test_scan_status_lists_every_computer_so_failures_are_not_silent(workbook):
    _, wb = workbook
    ws = wb["Scan Status"]
    by_name = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert [c.value for c in by_name["PRINTER-2"]][1:] == ["Failed", 0, "Access is denied."]
    assert by_name["PC-C"][1].value == "Partial" and by_name["PC-A"][1].value == "OK"
    colours = {n: r[1].fill.start_color.rgb[-6:] for n, r in by_name.items()}
    assert colours == {"PC-A": "C8E6C9", "PC-B": "C8E6C9", "PC-C": "FFF9C4", "PRINTER-2": "FFCDD2"}


def test_software_export_blocks_formula_injection_on_every_sheet(workbook):
    path, wb = workbook
    zf = zipfile.ZipFile(path)
    for member in zf.namelist():
        if member.startswith("xl/worksheets/sheet"):
            assert "<f>" not in zf.read(member).decode(), member
    titles = {r[1] for r in wb["Software"].iter_rows(min_row=2, values_only=True)}
    assert '=HYPERLINK("http://evil.example","x")' in titles


def test_software_export_of_an_empty_scan(tmp_path):
    path = tmp_path / "empty.xlsx"
    export_excel.export_software(str(path), [])
    assert load_workbook(path).sheetnames == ["Software", "Summary", "Scan Status"]
